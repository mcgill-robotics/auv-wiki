#!/usr/bin/env python3
"""Convert DEV_GUIDE.md into a formatted Word (.docx) document.

Handles the subset of Markdown used in the guide: ATX headings, fenced code
blocks, GitHub-style tables, blockquotes, ordered/unordered lists, horizontal
rules, and inline `code`/**bold** spans.
"""
import re
import sys

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


SRC = "timesheet-validator/DEV_GUIDE.md"
OUT = "timesheet-validator/DEV_GUIDE.docx"

CODE_FONT = "Consolas"
CODE_SHADE = "F2F2F2"


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_code_block(doc, lines):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Inches(0.15)
    para.paragraph_format.space_before = Pt(4)
    para.paragraph_format.space_after = Pt(8)
    # light grey background via paragraph shading
    p_pr = para._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), CODE_SHADE)
    p_pr.append(shd)
    run = para.add_run("\n".join(lines))
    run.font.name = CODE_FONT
    run.font.size = Pt(9)


def add_inline(paragraph, text):
    """Render inline **bold** and `code` spans into runs."""
    for piece in re.split(r"(\*\*.+?\*\*|`[^`]+`)", text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            run = paragraph.add_run(piece[2:-2])
            run.bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = CODE_FONT
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
        else:
            # strip leftover markdown links -> keep text
            piece = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", piece)
            paragraph.add_run(piece)


def parse_table_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, cell_text in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.paragraphs[0].text = ""
        run = cell.paragraphs[0].add_run(cell_text)
        run.bold = True
        shade_cell(cell, "D9E2F3")
    for row in body:
        cells = table.add_row().cells
        for i, cell_text in enumerate(row):
            if i >= len(cells):
                break
            cells[i].paragraphs[0].text = ""
            add_inline(cells[i].paragraphs[0], cell_text)
    doc.add_paragraph()


def main():
    with open(SRC, encoding="utf-8") as f:
        lines = f.read().split("\n")

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]

        # Fenced code block
        if line.startswith("```"):
            block = []
            i += 1
            while i < n and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            add_code_block(doc, block)
            i += 1
            continue

        # Table (line with |, next line is separator)
        if line.strip().startswith("|") and i + 1 < n and re.match(
            r"^\s*\|?[\s:|-]+\|?\s*$", lines[i + 1]
        ):
            rows = [parse_table_row(line)]
            i += 2  # skip header + separator
            while i < n and lines[i].strip().startswith("|"):
                rows.append(parse_table_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue

        # Headings
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1:
                h = doc.add_heading("", level=0)
                add_inline(h, text)
            else:
                h = doc.add_heading("", level=min(level - 1, 4) or 1)
                add_inline(h, text)
            i += 1
            continue

        # Horizontal rule
        if re.match(r"^---+\s*$", line):
            doc.add_paragraph().add_run("_" * 60).font.color.rgb = RGBColor(
                0xBB, 0xBB, 0xBB
            )
            i += 1
            continue

        # Blockquote
        if line.startswith(">"):
            quote_lines = []
            while i < n and lines[i].startswith(">"):
                quote_lines.append(lines[i].lstrip(">").strip())
                i += 1
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Inches(0.3)
            add_inline(para, " ".join(quote_lines))
            for run in para.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            continue

        # Ordered list
        m = re.match(r"^(\d+)\.\s+(.*)$", line)
        if m:
            para = doc.add_paragraph(style="List Number")
            add_inline(para, m.group(2))
            i += 1
            continue

        # Unordered list
        m = re.match(r"^[-*]\s+(.*)$", line)
        if m:
            para = doc.add_paragraph(style="List Bullet")
            add_inline(para, m.group(1))
            i += 1
            continue

        # Blank line
        if not line.strip():
            i += 1
            continue

        # Normal paragraph
        para = doc.add_paragraph()
        add_inline(para, line)
        i += 1

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    sys.exit(main())
